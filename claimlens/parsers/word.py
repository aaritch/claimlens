from __future__ import annotations

import io
import zipfile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from claimlens.parsers.base import ParsedDocument, ParseError


def parse_docx(data: bytes, filename: str) -> ParsedDocument:
    try:
        doc = Document(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ParseError(f"{filename}: unreadable DOCX ({exc.__class__.__name__})") from exc

    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # tables hold a lot of structured info in real forms
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return ParsedDocument(filename=filename, file_type="docx", text="\n".join(lines))
